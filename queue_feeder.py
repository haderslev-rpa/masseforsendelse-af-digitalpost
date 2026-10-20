"""Manuel queue-feeder til masseforsendelse-af-digitalpost.

Feederen læser en JSON-fil fra mappen forsendelser/, validerer SharePoint-
filerne og opretter ét work item pr. CPR-nummer.
"""

from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
import re
from typing import Any

from automation_server_client import AutomationServer
from openpyxl import load_workbook

from configuration import (
    SHAREPOINT_BASE_FOLDER,
    SHAREPOINT_LIBRARY_NAME,
    SHAREPOINT_SITE_NAME,
)
from sharepoint_template_service import (
    build_relative_path,
    download_file_by_drive_ids,
    get_file_metadata_by_path,
)

PROJECT_ROOT = Path(__file__).resolve().parent
SHIPMENT_CONFIG_DIRECTORY = PROJECT_ROOT / "forsendelser"


def _require_dict(value: Any, field_name: str) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise ValueError(f"{field_name} skal være et JSON-object.")
    return value


def _require_text(value: Any, field_name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field_name} skal udfyldes med tekst.")
    return value.strip()


def _choose_configuration_file() -> Path:
    """Spørg efter JSON-filens navn og returnér filens lokale sti."""
    SHIPMENT_CONFIG_DIRECTORY.mkdir(exist_ok=True)
    available_files = sorted(SHIPMENT_CONFIG_DIRECTORY.glob("*.json"))

    print("\nTilgængelige forsendelsesfiler:")
    for file_path in available_files:
        print(f"  - {file_path.name}")

    while True:
        entered_name = input(
            "\nSkriv navnet på JSON-filen, der skal bruges: "
        ).strip()

        if not entered_name:
            print("Du skal skrive et filnavn.")
            continue

        if (
            entered_name.startswith("source ")
            or "bin/activate" in entered_name
            or "debugpy" in entered_name
        ):
            print("En VS Code-kommando blev ignoreret. Skriv filnavnet igen.")
            continue

        if Path(entered_name).name != entered_name:
            print("Skriv kun filnavnet uden mapper eller fuld sti.")
            continue

        if not entered_name.lower().endswith(".json"):
            entered_name += ".json"

        configuration_path = SHIPMENT_CONFIG_DIRECTORY / entered_name
        if not configuration_path.is_file():
            print(f"Filen blev ikke fundet: {configuration_path}")
            continue

        return configuration_path


def _load_configuration(configuration_path: Path) -> dict[str, Any]:
    """Læs og returnér udsendelsens JSON-konfiguration."""
    try:
        with configuration_path.open("r", encoding="utf-8") as file:
            value = json.load(file)
    except json.JSONDecodeError as error:
        raise ValueError(
            f"JSON-fejl i {configuration_path.name}, linje {error.lineno}, "
            f"kolonne {error.colno}: {error.msg}"
        ) from error
    return _require_dict(value, "Forsendelseskonfigurationen")


def _validate_configuration(raw: dict[str, Any]) -> dict[str, Any]:
    """Validér JSON og returnér de normaliserede værdier."""
    excel = _require_dict(raw.get("excel"), "excel")
    document = _require_dict(raw.get("document"), "document")
    word = _require_dict(raw.get("word"), "word")
    attachments = _require_dict(raw.get("attachments"), "attachments")

    batch_id = _require_text(raw.get("batch_id"), "batch_id")
    if ":" in batch_id:
        raise ValueError("batch_id må ikke indeholde kolon (:).")

    mode = _require_text(document.get("mode"), "document.mode").upper()
    if mode not in {"WORD", "PDF"}:
        raise ValueError("document.mode skal være WORD eller PDF.")

    document_file_name = _require_text(
        document.get("file_name"),
        "document.file_name",
    )
    expected_extension = ".docx" if mode == "WORD" else ".pdf"
    if not document_file_name.lower().endswith(expected_extension):
        raise ValueError(
            f"document.file_name skal ende på {expected_extension}."
        )

    output_file_name = _require_text(
        document.get("output_file_name"),
        "document.output_file_name",
    )
    if not output_file_name.lower().endswith(".pdf"):
        raise ValueError("document.output_file_name skal ende på .pdf.")

    excel_file_name = _require_text(excel.get("file_name"), "excel.file_name")
    if not excel_file_name.lower().endswith(".xlsx"):
        raise ValueError("excel.file_name skal ende på .xlsx.")

    first_data_row = excel.get("first_data_row")
    if isinstance(first_data_row, bool) or not isinstance(first_data_row, int):
        raise ValueError("excel.first_data_row skal være et helt tal.")

    sheet_name = excel.get("sheet_name")
    if sheet_name is not None and not isinstance(sheet_name, str):
        raise ValueError("excel.sheet_name skal være tekst eller null.")

    fill_name_and_address = word.get("fill_name_and_address_bookmarks")
    if not isinstance(fill_name_and_address, bool):
        raise ValueError(
            "word.fill_name_and_address_bookmarks skal være true eller false."
        )

    bookmark_values = _require_dict(
        word.get("bookmark_values"),
        "word.bookmark_values",
    )

    delivery_method = _require_text(
        raw.get("delivery_method"),
        "delivery_method",
    ).upper()
    if delivery_method not in {
        "ONLY_DIGITAL_POST",
        "DIGITAL_OR_PHYSICAL_POST",
    }:
        raise ValueError("delivery_method har en ukendt værdi.")
    if mode == "PDF" and delivery_method == "DIGITAL_OR_PHYSICAL_POST":
        raise ValueError("PDF-mode kan kun bruges med ONLY_DIGITAL_POST.")
    if (
        mode == "WORD"
        and delivery_method == "DIGITAL_OR_PHYSICAL_POST"
        and not fill_name_and_address
    ):
        raise ValueError(
            "DIGITAL_OR_PHYSICAL_POST kræver automatisk navn og adresse."
        )

    dry_run = raw.get("dry_run")
    if not isinstance(dry_run, bool):
        raise ValueError("dry_run skal være true eller false.")

    validated_attachments: dict[str, str] = {}
    for display_name, sharepoint_name in attachments.items():
        display_name = _require_text(display_name, "attachments-nøgle")
        sharepoint_name = _require_text(
            sharepoint_name,
            f"attachments[{display_name!r}]",
        )
        if not display_name.lower().endswith(".pdf"):
            raise ValueError(f"Bilagsnavnet skal ende på .pdf: {display_name}")
        if not sharepoint_name.lower().endswith(".pdf"):
            raise ValueError(
                f"SharePoint-bilaget skal ende på .pdf: {sharepoint_name}"
            )
        validated_attachments[display_name] = sharepoint_name

    return {
        "batch_id": batch_id,
        "sharepoint_folder_name": _require_text(
            raw.get("sharepoint_folder_name"),
            "sharepoint_folder_name",
        ),
        "excel_file_name": excel_file_name,
        "excel_sheet_name": sheet_name,
        "cpr_column": _require_text(
            excel.get("cpr_column"),
            "excel.cpr_column",
        ).upper(),
        "first_data_row": first_data_row,
        "document_mode": mode,
        "document_file_name": document_file_name,
        "output_file_name": output_file_name,
        "subject": _require_text(raw.get("subject"), "subject"),
        "bookmark_values": bookmark_values if mode == "WORD" else {},
        "fill_name_and_address": (
            fill_name_and_address if mode == "WORD" else False
        ),
        "attachments": validated_attachments,
        "delivery_method": delivery_method,
        "dry_run": dry_run,
    }


def _find_sharepoint_file(
    configuration: dict[str, Any],
    file_name: str,
) -> dict[str, str]:
    """Find en fil og returnér kun metadata, som worker-processen behøver."""
    relative_path = build_relative_path(
        base_folder=SHAREPOINT_BASE_FOLDER,
        template_folder=configuration["sharepoint_folder_name"],
        file_name=file_name,
    )
    metadata = get_file_metadata_by_path(
        site_name=SHAREPOINT_SITE_NAME,
        library_name=SHAREPOINT_LIBRARY_NAME,
        relative_path=relative_path,
    )
    return {
        "drive_id": metadata["drive_id"],
        "drive_item_id": metadata["drive_item_id"],
        "file_name": metadata["file_name"],
    }


def _download_file(metadata: dict[str, str]) -> bytes:
    """Download en fil ud fra de minimale SharePoint-ID'er."""
    return download_file_by_drive_ids(
        drive_id=metadata["drive_id"],
        drive_item_id=metadata["drive_item_id"],
    )


def _normalize_cpr(value: object) -> str:
    if value is None:
        return ""
    text = str(value).zfill(10) if isinstance(value, int) else str(value).strip()
    return re.sub(r"\D", "", text)


def _read_cpr_numbers(
    excel_bytes: bytes,
    configuration: dict[str, Any],
) -> list[str]:
    """Læs, validér og returnér alle unikke CPR-numre."""
    workbook = load_workbook(
        BytesIO(excel_bytes),
        read_only=True,
        data_only=True,
    )
    try:
        sheet_name = configuration["excel_sheet_name"]
        worksheet = workbook[sheet_name] if sheet_name else workbook.active
        cpr_numbers: list[str] = []
        errors: list[str] = []
        seen: set[str] = set()

        for row_number in range(
            configuration["first_data_row"],
            worksheet.max_row + 1,
        ):
            value = worksheet[
                f"{configuration['cpr_column']}{row_number}"
            ].value
            cpr = _normalize_cpr(value)
            if not cpr:
                continue
            if len(cpr) != 10:
                errors.append(f"Række {row_number}: CPR skal have 10 cifre.")
            elif cpr in seen:
                errors.append(f"Række {row_number}: CPR forekommer flere gange.")
            else:
                seen.add(cpr)
                cpr_numbers.append(cpr)

        if errors:
            raise ValueError("Excel-filen indeholder fejl:\n" + "\n".join(errors))
        if not cpr_numbers:
            raise ValueError("Excel-filen indeholder ingen CPR-numre.")
        return cpr_numbers
    finally:
        workbook.close()


def _resolve_attachments(
    configuration: dict[str, Any],
) -> list[dict[str, str]]:
    """Find og kontrollér fælles PDF-bilag."""
    result: list[dict[str, str]] = []
    for display_name, sharepoint_name in configuration["attachments"].items():
        metadata = _find_sharepoint_file(configuration, sharepoint_name)
        if not _download_file(metadata).startswith(b"%PDF-"):
            raise ValueError(f"Bilaget {sharepoint_name!r} er ikke en gyldig PDF.")
        result.append(
            {
                "drive_id": metadata["drive_id"],
                "drive_item_id": metadata["drive_item_id"],
                "file_name": display_name,
            }
        )
    return result


def main() -> None:
    configuration_path = _choose_configuration_file()
    configuration = _validate_configuration(
        _load_configuration(configuration_path)
    )
    automation_server = AutomationServer.from_environment()

    excel_file = _find_sharepoint_file(
        configuration,
        configuration["excel_file_name"],
    )
    main_document = _find_sharepoint_file(
        configuration,
        configuration["document_file_name"],
    )
    attachments = _resolve_attachments(configuration)
    cpr_numbers = _read_cpr_numbers(
        _download_file(excel_file),
        configuration,
    )

    if configuration["document_mode"] == "PDF":
        if not _download_file(main_document).startswith(b"%PDF-"):
            raise ValueError("Hoveddokumentet er ikke en gyldig PDF-fil.")

    print("\n" + "=" * 70)
    print("KONTROL AF MASSEFORSENDELSE")
    print("=" * 70)
    print(f"Konfigurationsfil: {configuration_path.name}")
    print(f"Batch-ID: {configuration['batch_id']}")
    print(f"Dokumenttype: {configuration['document_mode']}")
    print(f"Antal CPR-numre: {len(cpr_numbers)}")
    print(f"Antal bilag: {len(attachments)}")
    print(f"Leveringsmetode: {configuration['delivery_method']}")

    if configuration["dry_run"]:
        print("\ndry_run=true. Der er ikke oprettet work items.")
        return

    workqueue = automation_server.workqueue()

    for sequence, cpr in enumerate(cpr_numbers, start=1):
        reference = f"{configuration['batch_id']}:{sequence:06d}:{cpr}"
        item_data = {
            "box": {
                "cpr": cpr,
                "subject": configuration["subject"],
                "shipment_configuration_file": configuration_path.name,
                "document_mode": configuration["document_mode"],
                "main_document": dict(main_document),
                "output_file_name": configuration["output_file_name"],
                "bookmark_values": dict(configuration["bookmark_values"]),
                "fill_name_and_address": configuration[
                    "fill_name_and_address"
                ],
                "attachments": [dict(value) for value in attachments],
                "delivery_method": configuration["delivery_method"],
                "processing_state": "READY",
                "digitalpost_result": None,
                "last_error": None,
            },
            "defer": None,
            "state": [],
            "status": {
                "status": "New",
                "status_code": "Klar til dokumentbehandling",
            },
        }
        workqueue.add_item(data=item_data, reference=reference)

        if sequence % 100 == 0:
            print(f"Oprettet: {sequence}/{len(cpr_numbers)}")

    print(f"\nFærdig. Oprettede {len(cpr_numbers)} work items.")


if __name__ == "__main__":
    main()
